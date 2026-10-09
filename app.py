"""
Asistente de Entrevistas Estructuradas por Competencias (AEEC)
---------------------------------------------------------------
Genera, a partir de la descripción de un puesto, un guion de entrevista
estructurada por competencias con escala BARS 1-5 y una hoja de puntuación
ponderada. La herramienta NO evalúa candidatos ni trata datos personales:
la valoración la realiza siempre una persona.

Ejecución local:  streamlit run app.py
"""

import io
import json
import re

import pandas as pd
import streamlit as st

# ------------------------------------------------------------------
# Configuración general
# ------------------------------------------------------------------
st.set_page_config(
    page_title="AEEC · Entrevistas por competencias",
    page_icon="🧭",
    layout="wide",
)

CATALOGO_COMPETENCIAS = [
    "Trabajo en equipo",
    "Orientación a resultados",
    "Comunicación",
    "Planificación y organización",
    "Orientación al cliente",
    "Adaptabilidad",
    "Resolución de problemas",
    "Liderazgo",
    "Iniciativa",
    "Negociación",
    "Compromiso con la calidad",
    "Orientación a la seguridad y prevención",
]
NIVELES = ["Básico", "Intermedio", "Avanzado"]

# Patrones para detectar datos personales que no deberían introducirse
PATRONES_DATOS_PERSONALES = {
    "DNI/NIE": r"\b[XYZ]?\d{7,8}[A-Z]\b",
    "correo electrónico": r"[\w\.-]+@[\w\.-]+\.\w{2,}",
    "teléfono": r"\b(?:\+34\s?)?[6789]\d{2}\s?\d{3}\s?\d{3}\b",
}


# ------------------------------------------------------------------
# Utilidades
# ------------------------------------------------------------------
def leer_secreto(clave, defecto=None):
    """Lee un valor de st.secrets sin romper la app si no existe el archivo."""
    try:
        return st.secrets.get(clave, defecto)
    except Exception:
        return defecto


def configuracion_ia():
    """Devuelve (proveedor, modelo, clave) o None si no hay clave → modo demo."""
    proveedor = str(leer_secreto("PROVEEDOR", "anthropic")).lower()
    modelo = leer_secreto("MODELO")
    clave = leer_secreto(
        "ANTHROPIC_API_KEY" if proveedor == "anthropic" else "OPENAI_API_KEY"
    )
    if not clave or not modelo:
        return None
    return proveedor, modelo, clave


def extraer_texto(archivo):
    """Extrae texto de un PDF, DOCX o TXT subido por el usuario."""
    nombre = archivo.name.lower()
    if nombre.endswith(".pdf"):
        from pypdf import PdfReader

        lector = PdfReader(archivo)
        return "\n".join((p.extract_text() or "") for p in lector.pages)
    if nombre.endswith(".docx"):
        from docx import Document

        doc = Document(archivo)
        return "\n".join(p.text for p in doc.paragraphs)
    return archivo.read().decode("utf-8", errors="ignore")


def detectar_datos_personales(texto):
    """Devuelve la lista de tipos de dato personal encontrados en el texto."""
    return [
        tipo
        for tipo, patron in PATRONES_DATOS_PERSONALES.items()
        if re.search(patron, texto or "", flags=re.IGNORECASE)
    ]


def limpiar_json(texto):
    """Quita bloques ``` y texto sobrante antes de json.loads."""
    texto = re.sub(r"```(?:json)?", "", texto).strip()
    inicio, fin = texto.find("{"), texto.rfind("}")
    return texto[inicio : fin + 1] if inicio != -1 and fin != -1 else texto


# ------------------------------------------------------------------
# Instrucciones al modelo
# ------------------------------------------------------------------
PROMPT_SISTEMA = """Eres un psicólogo del trabajo experto en selección por competencias
y entrevista conductual estructurada. Trabajas para el departamento de Personas de
Nexalia Logística, S. L., empresa privada española de distribución y almacenaje
(1.200 empleados, 9 centros). Las entrevistas las realizan técnicos de selección y
mandos de centro, por lo que el guion debe ser claro y fácil de aplicar.
Reglas:
- Responde SOLO con un objeto JSON válido, sin texto adicional ni ```.
- Escribe en español de España, con lenguaje inclusivo y claro.
- Las preguntas deben referirse a conductas y situaciones laborales, nunca a la vida
  personal, salud, estado civil, maternidad/paternidad, religión, origen, ideología,
  afiliación sindical, orientación sexual ni edad (Ley 15/2022).
- Las anclas BARS describen conductas observables, de 1 (muy por debajo) a 5 (excelente).
Estructura obligatoria:
{"apertura": str,
 "competencias": [{"nombre": str, "definicion": str, "indicadores": [str],
   "preguntas": [{"pregunta": str, "tipo": str, "repreguntas": [str]}],
   "bars": {"1": str, "2": str, "3": str, "4": str, "5": str},
   "alertas": [str]}],
 "cierre": str,
 "preguntas_prohibidas": [str],
 "control_sesgos": [str]}"""


def construir_prompt_usuario(datos, competencias=None):
    comps = competencias or datos["competencias"]
    lineas = [
        f"- {c['Competencia']} (peso {c['Peso (%)']} %, nivel {c['Nivel']})"
        for c in comps
    ]
    return f"""Puesto: {datos['puesto']} | Área: {datos['area']}
Duración de la entrevista: {datos['duracion']} minutos
Tipo de preguntas: {datos['tipo']} | Preguntas por competencia: {datos['n_preguntas']}
Valores o contexto corporativo: {datos['valores'] or 'no indicados'}

Descripción del puesto:
{datos['descripcion'][:8000]}

Competencias a evaluar:
{chr(10).join(lineas)}

Genera el guion completo. Incluye 3-4 indicadores por competencia, 2 repreguntas por
pregunta, 2-3 señales de alerta, 6-8 preguntas prohibidas y 5-6 puntos de control de sesgos."""


def llamar_modelo(prompt_usuario, cfg):
    """Llama al proveedor configurado y devuelve el texto de respuesta."""
    proveedor, modelo, clave = cfg
    if proveedor == "openai":
        from openai import OpenAI

        cliente = OpenAI(api_key=clave)
        resp = cliente.chat.completions.create(
            model=modelo,
            response_format={"type": "json_object"},
            messages=[
                {"role": "system", "content": PROMPT_SISTEMA},
                {"role": "user", "content": prompt_usuario},
            ],
        )
        return resp.choices[0].message.content
    import anthropic

    cliente = anthropic.Anthropic(api_key=clave)
    resp = cliente.messages.create(
        model=modelo,
        max_tokens=8000,
        system=PROMPT_SISTEMA,
        messages=[{"role": "user", "content": prompt_usuario}],
    )
    return "".join(b.text for b in resp.content if b.type == "text")


def generar_guion(datos, cfg, competencias=None):
    """Genera el guion con un reintento si el JSON no es válido."""
    prompt = construir_prompt_usuario(datos, competencias)
    ultimo_error = None
    for _ in range(2):
        try:
            return json.loads(limpiar_json(llamar_modelo(prompt, cfg)))
        except json.JSONDecodeError as e:
            ultimo_error = e
    raise ValueError(f"La IA no devolvió un formato válido ({ultimo_error}).")


def guion_demo(datos):
    """Ejemplo precargado para usar la app sin clave de API."""
    comps = []
    for c in datos["competencias"]:
        n = c["Competencia"]
        comps.append(
            {
                "nombre": n,
                "definicion": f"Capacidad de demostrar {n.lower()} de forma consistente en el desempeño del puesto.",
                "indicadores": [
                    "Aporta ejemplos concretos y recientes",
                    "Describe su papel personal, no solo el del equipo",
                    "Explica resultados medibles",
                ],
                "preguntas": [
                    {
                        "pregunta": f"Cuénteme una situación reciente en la que tuvo que poner en práctica {n.lower()}. ¿Qué hizo exactamente?",
                        "tipo": "Conductual (STAR)",
                        "repreguntas": [
                            "¿Cuál fue el resultado?",
                            "¿Qué haría distinto hoy?",
                        ],
                    }
                ],
                "bars": {
                    "1": "No aporta ejemplos o son genéricos.",
                    "2": "Ejemplo poco claro, papel personal difuso.",
                    "3": "Ejemplo adecuado con acciones concretas.",
                    "4": "Ejemplo sólido, con resultados y reflexión.",
                    "5": "Varios ejemplos de alto impacto, transferibles al puesto.",
                },
                "alertas": ["Respuestas hipotéticas en lugar de hechos", "Atribuye todo al equipo"],
            }
        )
    return {
        "apertura": "Presentación del comité, objetivo de la entrevista, duración y aviso de toma de notas.",
        "competencias": comps,
        "cierre": "Turno de preguntas de la persona candidata y explicación de los siguientes pasos.",
        "preguntas_prohibidas": [
            "Estado civil o planes familiares",
            "Embarazo o maternidad/paternidad",
            "Religión o creencias",
            "Origen o nacionalidad",
            "Estado de salud",
            "Afiliación sindical o política",
        ],
        "control_sesgos": [
            "Mismas preguntas y orden para todas las personas candidatas",
            "Puntuar cada competencia justo al terminarla",
            "Evitar el efecto halo",
            "Basar la nota en conductas, no en impresiones",
            "Puntuar de forma independiente antes de deliberar",
        ],
    }


# ------------------------------------------------------------------
# Exportación
# ------------------------------------------------------------------
def guion_a_docx(guion, datos):
    from docx import Document

    doc = Document()
    doc.add_heading(f"Guion de entrevista – {datos['puesto']}", 0)
    doc.add_paragraph(f"Área: {datos['area']} · Duración: {datos['duracion']} min")
    doc.add_paragraph().add_run("Contenido generado con apoyo de IA y revisado por el departamento de Personas.").italic = True
    doc.add_heading("Apertura", 1)
    doc.add_paragraph(guion.get("apertura", ""))
    for c in guion.get("competencias", []):
        doc.add_heading(c["nombre"], 1)
        doc.add_paragraph(c.get("definicion", ""))
        doc.add_heading("Indicadores", 2)
        for i in c.get("indicadores", []):
            doc.add_paragraph(i, style="List Bullet")
        doc.add_heading("Preguntas", 2)
        for p in c.get("preguntas", []):
            doc.add_paragraph(f"{p['pregunta']} ({p.get('tipo', '')})", style="List Number")
            for r in p.get("repreguntas", []):
                doc.add_paragraph(f"→ {r}")
        doc.add_heading("Escala BARS", 2)
        tabla = doc.add_table(rows=1, cols=2)
        tabla.style = "Table Grid"
        tabla.rows[0].cells[0].text, tabla.rows[0].cells[1].text = "Nivel", "Conducta"
        for nivel, desc in c.get("bars", {}).items():
            fila = tabla.add_row().cells
            fila[0].text, fila[1].text = str(nivel), desc
        doc.add_heading("Señales de alerta", 2)
        for a in c.get("alertas", []):
            doc.add_paragraph(a, style="List Bullet")
    doc.add_heading("Cierre", 1)
    doc.add_paragraph(guion.get("cierre", ""))
    doc.add_heading("Preguntas que no deben formularse", 1)
    for p in guion.get("preguntas_prohibidas", []):
        doc.add_paragraph(p, style="List Bullet")
    doc.add_heading("Control de sesgos", 1)
    for p in guion.get("control_sesgos", []):
        doc.add_paragraph(p, style="List Bullet")
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


# ------------------------------------------------------------------
# Interfaz
# ------------------------------------------------------------------
cfg = configuracion_ia()

with st.sidebar:
    st.header("⚙️ Configuración")
    tipo = st.selectbox("Tipo de pregunta", ["Conductual (STAR)", "Situacional", "Mixta"])
    n_preguntas = st.slider("Preguntas por competencia", 2, 3, 2)
    duracion = st.selectbox("Duración (min)", [30, 45, 60], index=1)
    st.divider()
    if cfg:
        st.success(f"IA activa: {cfg[0]} · {cfg[1]}")
    else:
        st.info("Modo demostración (sin clave de API). Se muestra un ejemplo.")
    st.warning("Contenido generado por IA: revise y adapte antes de usar. La valoración de candidatos es siempre humana.")

st.title("🧭 Asistente de Entrevistas Estructuradas por Competencias")
st.caption("Nexalia Logística, S. L. (empresa ficticia) · Guiones de entrevista con escala BARS a partir de la descripción del puesto.")

tab1, tab2, tab3 = st.tabs(["1. Puesto y competencias", "2. Guion generado", "3. Puntuación y exportación"])

# --- Pestaña 1 -----------------------------------------------------
with tab1:
    c1, c2 = st.columns(2)
    puesto = c1.text_input("Denominación del puesto *")
    area = c2.text_input("Área / departamento *")
    descripcion = st.text_area("Descripción y funciones del puesto *", height=180)
    archivo = st.file_uploader("…o suba la descripción (PDF, DOCX, TXT)", type=["pdf", "docx", "txt"])
    if archivo:
        descripcion = (descripcion + "\n" + extraer_texto(archivo)).strip()
        st.caption(f"Texto extraído: {len(descripcion)} caracteres.")

    elegidas = st.multiselect("Competencias del catálogo (3–6)", CATALOGO_COMPETENCIAS)
    libre = st.text_input("Competencia adicional (opcional)")
    lista = elegidas + ([libre] if libre.strip() else [])

    competencias = []
    if lista:
        peso_base = round(100 / len(lista))
        df = pd.DataFrame(
            {
                "Competencia": lista,
                "Peso (%)": [peso_base] * (len(lista) - 1) + [100 - peso_base * (len(lista) - 1)],
                "Nivel": ["Intermedio"] * len(lista),
            }
        )
        df = st.data_editor(
            df,
            hide_index=True,
            width="stretch",
            disabled=["Competencia"],
            column_config={
                "Peso (%)": st.column_config.NumberColumn(min_value=0, max_value=100, step=5),
                "Nivel": st.column_config.SelectboxColumn(options=NIVELES),
            },
        )
        competencias = df.to_dict("records")
        st.caption(f"Suma de pesos: {int(df['Peso (%)'].sum())} %")

    valores = st.text_area("Valores corporativos o contexto (opcional)", height=80)

    if st.button("Generar guion", type="primary"):
        errores = []
        if not puesto or not area or not descripcion.strip():
            errores.append("Complete puesto, área y descripción.")
        if not 3 <= len(competencias) <= 6:
            errores.append("Seleccione entre 3 y 6 competencias.")
        elif int(sum(c["Peso (%)"] for c in competencias)) != 100:
            errores.append("Los pesos deben sumar 100 %.")
        for e in errores:
            st.error(e)

        hallazgos = detectar_datos_personales(descripcion + " " + valores)
        if hallazgos:
            st.warning(f"Posibles datos personales detectados ({', '.join(hallazgos)}). Elimínelos: la herramienta solo necesita datos del puesto.")

        if not errores:
            datos = dict(
                puesto=puesto, area=area, descripcion=descripcion, valores=valores,
                competencias=competencias, tipo=tipo, n_preguntas=n_preguntas, duracion=duracion,
            )
            with st.spinner("Generando guion…"):
                try:
                    guion = generar_guion(datos, cfg) if cfg else guion_demo(datos)
                    st.session_state["guion"], st.session_state["datos"] = guion, datos
                    st.success("Guion generado. Revíselo en la pestaña 2.")
                except Exception as e:
                    st.error(f"No se pudo generar el guion: {e}")

# --- Pestaña 2 -----------------------------------------------------
with tab2:
    guion = st.session_state.get("guion")
    if not guion:
        st.info("Genere primero un guion en la pestaña 1.")
    else:
        datos = st.session_state["datos"]
        st.subheader("Apertura")
        st.write(guion.get("apertura", ""))
        for idx, comp in enumerate(guion.get("competencias", [])):
            with st.expander(f"🔹 {comp['nombre']}", expanded=idx == 0):
                st.markdown(f"**Definición:** {comp.get('definicion', '')}")
                st.markdown("**Indicadores observables**")
                for i in comp.get("indicadores", []):
                    st.markdown(f"- {i}")
                st.markdown("**Preguntas**")
                for n, p in enumerate(comp.get("preguntas", []), 1):
                    st.markdown(f"{n}. {p['pregunta']} *({p.get('tipo', '')})*")
                    for r in p.get("repreguntas", []):
                        st.markdown(f"&emsp;↳ {r}")
                st.markdown("**Escala BARS**")
                st.table(pd.DataFrame(comp.get("bars", {}).items(), columns=["Nivel", "Conducta observable"]))
                st.markdown("**Señales de alerta:** " + " · ".join(comp.get("alertas", [])))
                if cfg and st.button("Regenerar esta competencia", key=f"regen_{idx}"):
                    origen = [c for c in datos["competencias"] if c["Competencia"] == comp["nombre"]]
                    with st.spinner("Regenerando…"):
                        try:
                            nuevo = generar_guion(datos, cfg, origen or None)
                            st.session_state["guion"]["competencias"][idx] = nuevo["competencias"][0]
                            st.rerun()
                        except Exception as e:
                            st.error(f"No se pudo regenerar: {e}")
        st.subheader("Cierre")
        st.write(guion.get("cierre", ""))
        c1, c2 = st.columns(2)
        with c1:
            st.markdown("**⛔ Preguntas que no deben formularse**")
            for p in guion.get("preguntas_prohibidas", []):
                st.markdown(f"- {p}")
        with c2:
            st.markdown("**✅ Control de sesgos**")
            for p in guion.get("control_sesgos", []):
                st.markdown(f"- {p}")

# --- Pestaña 3 -----------------------------------------------------
with tab3:
    guion = st.session_state.get("guion")
    if not guion:
        st.info("Genere primero un guion en la pestaña 1.")
    else:
        datos = st.session_state["datos"]
        codigo = st.text_input("Código de candidato/a (no use nombres)", value="CAND-001")
        filas = []
        for c in datos["competencias"]:
            col1, col2 = st.columns([1, 2])
            nota = col1.slider(c["Competencia"], 1, 5, 3, key=f"nota_{c['Competencia']}")
            obs = col2.text_input("Observaciones", key=f"obs_{c['Competencia']}")
            filas.append(
                {"Código": codigo, "Competencia": c["Competencia"], "Peso (%)": c["Peso (%)"],
                 "Puntuación (1-5)": nota, "Observaciones": obs}
            )
        hoja = pd.DataFrame(filas)
        # Cálculo local, sin IA
        media = (hoja["Puntuación (1-5)"] * hoja["Peso (%)"]).sum() / hoja["Peso (%)"].sum()
        st.metric("Puntuación ponderada", f"{media:.2f} / 5")
        st.bar_chart(hoja.set_index("Competencia")["Puntuación (1-5)"])

        c1, c2 = st.columns(2)
        c1.download_button(
            "📄 Descargar guion (.docx)",
            data=guion_a_docx(guion, datos),
            file_name=f"guion_{datos['puesto'].replace(' ', '_')}.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        c2.download_button(
            "📊 Descargar hoja de puntuación (.csv)",
            data=hoja.to_csv(index=False, sep=";").encode("utf-8-sig"),
            file_name=f"puntuacion_{codigo}.csv",
            mime="text/csv",
        )
